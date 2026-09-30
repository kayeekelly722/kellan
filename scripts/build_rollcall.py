#!/usr/bin/env python3
"""Build the 船河樂2026 boarding roll-call workbook."""

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.workbook.properties import CalcProperties
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.page import PageMargins

# Group name, headcount. Order matches the sign-up list.
GROUPS = [
    ("高教練", 8),
    ("莫朗森", 2),
    ("賴俊瑋", 2),
    ("王知淳", 2),
    ("陳彥生", 3),
    ("王光華", 1),
    ("簡珈晴", 2),
    ("吳承祐", 1),
    ("吳柏皜", 2),
    ("許詩雅", 2),
    ("李俊謙家長", 1),
    ("陳展芹家長", 1),
    ("黃思瀚", 3),
    ("呂卓衡", 3),
    ("潘芊庭", 2),
    ("唐奕忻", 2),
    ("洪嘉駿", 5),
    ("李志斌", 3),
    ("葉城佑", 2),
    ("湯sir", 3),
    ("祈銘勛", 1),
]

MAX_SLOTS = 8
EXPECTED_TOTAL = 51
FIRST_DATA_ROW = 7
LAST_DATA_ROW = FIRST_DATA_ROW + len(GROUPS) - 1  # 27
TOTAL_ROW = LAST_DATA_ROW + 1  # 28

CHECK = "✓"
BOX = "□"
NA = "–"
CIRCLED = ["①", "②", "③", "④", "⑤", "⑥", "⑦", "⑧"]

NAVY = "1B4F72"
NAVY_DEEP = "154360"
INK = "1C2833"
WHITE = "FFFFFF"
TEAL = "0E6655"
SEA = "148F77"
SAND = "FBF6EA"
PAPER = "F7FBFD"
LINE = "D5DDE5"
GOLD = "C59B2D"
GOLD_BG = "FFF4CC"
GOLD_HEADER = "B7950B"
GREY_BG = "EEF1F4"
GREY_FONT = "B0B8C1"
GREEN_BG = "E8F8F0"
GREEN_FONT = "0E6655"
AMBER_BG = "FEF6DD"
AMBER_FONT = "9A6B12"
RED_BG = "FDEDEC"
RED_FONT = "922B21"
SLATE = "5D6D7E"
BLUE_SOFT = "EAF2F8"
OK_GREEN = "196F3D"

FONT_NAME = "Microsoft JhengHei"
TICK_FONT = "Segoe UI Symbol"


def font(size=12, bold=False, color=INK, name=FONT_NAME):
    return Font(name=name, size=size, bold=bold, color=color)


def fill(hex_color):
    return PatternFill("solid", fgColor=hex_color)


def side(style="thin", color=LINE):
    return Side(style=style, color=color)


def border(color=LINE, style="thin"):
    edge = side(style, color)
    return Border(left=edge, right=edge, top=edge, bottom=edge)


def apply(cell, value=None, font_=None, fill_=None, align=None, border_=None, locked=True, fmt=None):
    if value is not None:
        cell.value = value
    if font_ is not None:
        cell.font = font_
    if fill_ is not None:
        cell.fill = fill_
    if align is not None:
        cell.alignment = align
    if border_ is not None:
        cell.border = border_
    if fmt is not None:
        cell.number_format = fmt
    cell.protection = Protection(locked=locked)


def merge(ws, range_ref, value=None, font_=None, fill_=None, align=None, border_=None, locked=True):
    from openpyxl.utils import range_boundaries

    min_col, min_row, max_col, max_row = range_boundaries(range_ref)
    if not (min_col == max_col and min_row == max_row):
        ws.merge_cells(range_ref)
    for row in range(min_row, max_row + 1):
        for col in range(min_col, max_col + 1):
            cell = ws.cell(row, col)
            apply(
                cell,
                value=value if (row == min_row and col == min_col) else None,
                font_=font_,
                fill_=fill_,
                align=align,
                border_=border_,
                locked=locked,
            )
    return ws.cell(min_row, min_col)


def build():
    if len(GROUPS) != 21:
        raise SystemExit(f"Expected 21 groups, got {len(GROUPS)}")
    total = sum(count for _, count in GROUPS)
    if total != EXPECTED_TOTAL:
        raise SystemExit(f"Expected {EXPECTED_TOTAL} people, got {total}")
    if any(count < 1 or count > MAX_SLOTS for _, count in GROUPS):
        raise SystemExit("A group size is outside 1–8")

    wb = Workbook()
    wb.calculation = CalcProperties(calcMode="auto", fullCalcOnLoad=True, forceFullCalc=True)
    wb.properties.title = "船河樂2026 上船點名表"
    wb.properties.subject = "21 組、51 位"
    wb.properties.category = "點名"

    options = wb.active
    options.title = "選項"
    options["A1"] = CHECK
    options["A2"] = BOX
    options.sheet_state = "hidden"
    wb.defined_names.add(DefinedName(name="點名選項", attr_text="選項!$A$1:$A$2"))

    ws = wb.create_sheet("點名表", 0)
    extras = wb.create_sheet("臨時加人", 1)
    build_roll(ws)
    build_extras(extras)

    out = "/workspace/船河樂2026_點名表.xlsx"
    wb.save(out)
    return out


def build_roll(ws):
    ws.sheet_properties.tabColor = NAVY
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 120
    ws.sheet_view.view = "normal"
    ws.freeze_panes = "A7"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.page_setup.horizontalCentered = True
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins = PageMargins(
        left=0.4, right=0.4, top=0.35, bottom=0.4, header=0.12, footer=0.16
    )
    ws.print_options.horizontalCentered = True
    ws.print_title_rows = "1:6"
    ws.page_setup.scale = 100
    ws.oddFooter.left.text = "黃格剔 ✓    灰格不用填"
    ws.oddFooter.left.font = FONT_NAME
    ws.oddFooter.left.size = 9
    ws.oddFooter.center.text = "船河樂2026 上船點名表"
    ws.oddFooter.center.font = FONT_NAME
    ws.oddFooter.center.size = 9
    ws.oddFooter.right.text = "共 51 位"
    ws.oddFooter.right.font = FONT_NAME
    ws.oddFooter.right.size = 9
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1

    widths = {
        "A": 10,
        "B": 18,
        "C": 10,
        "L": 9,
        "M": 9,
        "N": 11,
        "O": 28,
    }
    for col in "DEFGHIJK":
        widths[col] = 5.6
    for col, width in widths.items():
        ws.column_dimensions[col].width = width

    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left = Alignment(horizontal="left", vertical="center", indent=1)
    right = Alignment(horizontal="right", vertical="center")

    ws.row_dimensions[1].height = 34
    ws.row_dimensions[2].height = 20
    ws.row_dimensions[3].height = 28
    ws.row_dimensions[4].height = 16
    ws.row_dimensions[5].height = 28
    ws.row_dimensions[6].height = 24

    merge(
        ws, "A1:O1",
        "船河樂2026    上船點名表",
        font(22, True, WHITE),
        fill(NAVY),
        center,
        border(NAVY),
    )
    merge(
        ws, "A2:O2",
        "共 21 組  ·  51 位　　黃格選 ✓（紙本直接在 □ 上打勾）　　灰格不用填　　已到／未到／狀態會自動計算",
        font(11, False, "EAF2F8"),
        fill(NAVY_DEEP),
        center,
        border(NAVY_DEEP),
    )

    label_fill = fill("F8F1DE")
    input_fill = fill("FFFFFF")
    input_border = border(NAVY, "medium")
    meta_border = border("F8F1DE")

    def meta_label(range_ref, text):
        merge(ws, range_ref, text, font(10, True, NAVY), label_fill, center, meta_border)

    def meta_input(range_ref):
        merge(
            ws,
            range_ref,
            None,
            font(12, False, INK),
            input_fill,
            Alignment(horizontal="left", vertical="center", indent=1),
            input_border,
            locked=False,
        )

    # Inputs sit in wider merges so 時間 and 負責人 can actually be written in.
    meta_label("A3:A3", "日期")
    meta_input("B3:C3")
    meta_label("D3:E3", "時間")
    meta_input("F3:H3")
    meta_label("I3:J3", "負責人")
    meta_input("K3:L3")
    meta_label("M3:M3", "地點")
    meta_input("N3:O3")

    # Live totals. Values sit in the top-left cell of each merged block.
    kpi_specs = [
        ("A4:C4", "A5:C5", "預計人數", f"=SUM(C{FIRST_DATA_ROW}:C{LAST_DATA_ROW})", NAVY, WHITE),
        ("D4:F4", "D5:F5", "已到", f"=SUM(L{FIRST_DATA_ROW}:L{LAST_DATA_ROW})", TEAL, WHITE),
        ("G4:I4", "G5:I5", "未到", f"=SUM(M{FIRST_DATA_ROW}:M{LAST_DATA_ROW})", "1A5276", WHITE),
        ("J4:L4", "J5:L5", "已齊組", f'=COUNTIF(N{FIRST_DATA_ROW}:N{LAST_DATA_ROW},"齊")&"/21"', SEA, WHITE),
        ("M4:O4", "M5:O5", "未齊組", f'=COUNTIF(N{FIRST_DATA_ROW}:N{LAST_DATA_ROW},"<>齊")', SLATE, WHITE),
    ]
    for label_range, value_range, label, formula, bg, fg in kpi_specs:
        merge(ws, label_range, label, font(9, True, fg), fill(bg), center, border(bg))
        merge(ws, value_range, formula, font(20, True, fg), fill(bg), center, border(bg))

    headers = ["編號", "姓名", "人數", *CIRCLED, "已到", "未到", "狀態", "備註"]
    header_fill = fill(NAVY)
    tick_header_fill = fill("1A6A7D")
    for col, text in enumerate(headers, 1):
        is_tick = 4 <= col <= 11
        apply(
            ws.cell(6, col),
            text,
            font(13 if is_tick else 11, True, WHITE, TICK_FONT if is_tick else FONT_NAME),
            tick_header_fill if is_tick else header_fill,
            center,
            border(NAVY),
        )
    ws.cell(6, 3).comment = Comment(
        "人數已鎖定，與黃格數目一致。\n名單以外的人請寫在「臨時加人」，不要改這裡。\n如要解鎖：校閱 → 取消保護工作表（無密碼）。",
        "點名表",
        width=240,
        height=80,
    )
    ws.cell(6, 4).comment = Comment(
        "每到一位，在黃格選 ✓。\n灰格代表該組沒有此人，不用填。\n紙本列印後，直接在 □ 上打勾即可。",
        "點名表",
        width=220,
        height=70,
    )

    thin = border()
    tick_border = border("E1C56A", "medium")
    na_border = border("E1E6EA")
    zebra = fill(PAPER)
    white = fill(WHITE)

    for index, (name, count) in enumerate(GROUPS):
        row = FIRST_DATA_ROW + index
        ws.row_dimensions[row].height = 22
        band = zebra if index % 2 else white

        apply(ws.cell(row, 1), index + 1, font(12, True, NAVY), band, center, thin)
        apply(ws.cell(row, 2), name, font(13, True, INK), band, left, thin, locked=False)
        apply(
            ws.cell(row, 3),
            count,
            font(12, True, INK),
            band,
            center,
            thin,
            fmt="0",
        )

        for slot in range(1, MAX_SLOTS + 1):
            cell = ws.cell(row, 3 + slot)
            if slot <= count:
                apply(
                    cell,
                    BOX,
                    font(16, False, "6B5A1E", TICK_FONT),
                    fill(GOLD_BG),
                    center,
                    tick_border,
                    locked=False,
                )
            else:
                apply(
                    cell,
                    NA,
                    font(12, False, GREY_FONT, TICK_FONT),
                    fill(GREY_BG),
                    center,
                    na_border,
                )

        apply(
            ws.cell(row, 12),
            f'=COUNTIF(D{row}:K{row},"{CHECK}")',
            font(12, True, TEAL),
            band,
            center,
            thin,
            fmt="0",
        )
        apply(
            ws.cell(row, 13),
            f"=C{row}-L{row}",
            font(12, False, INK),
            band,
            center,
            thin,
            fmt="0",
        )
        apply(
            ws.cell(row, 14),
            f'=IF(L{row}>C{row},"超出",IF(L{row}=C{row},"齊",IF(L{row}=0,"未到","未齊")))',
            font(11, True, SLATE),
            band,
            center,
            thin,
        )
        apply(ws.cell(row, 15), None, font(11, False, INK), band, left, thin, locked=False)

    dv = DataValidation(
        type="list",
        formula1="=點名選項",
        allow_blank=False,
        showDropDown=False,
        showErrorMessage=True,
        showInputMessage=False,
        errorTitle="點名",
        error="請選 ✓ 或 □",
        errorStyle="stop",
    )
    for index, (_, count) in enumerate(GROUPS):
        row = FIRST_DATA_ROW + index
        last_col = get_column_letter(3 + count)
        dv.add(f"D{row}:{last_col}{row}")
    ws.add_data_validation(dv)

    # Status colours. 未到 stays neutral so the sheet is calm before anyone arrives.
    ws.conditional_formatting.add(
        f"N{FIRST_DATA_ROW}:N{LAST_DATA_ROW}",
        CellIsRule(operator="equal", formula=['"齊"'], fill=fill(GREEN_BG), font=font(11, True, OK_GREEN)),
    )
    ws.conditional_formatting.add(
        f"N{FIRST_DATA_ROW}:N{LAST_DATA_ROW}",
        CellIsRule(operator="equal", formula=['"未齊"'], fill=fill(AMBER_BG), font=font(11, True, AMBER_FONT)),
    )
    ws.conditional_formatting.add(
        f"N{FIRST_DATA_ROW}:N{LAST_DATA_ROW}",
        CellIsRule(operator="equal", formula=['"超出"'], fill=fill(RED_BG), font=font(11, True, RED_FONT)),
    )
    ws.conditional_formatting.add(
        f"D{FIRST_DATA_ROW}:K{LAST_DATA_ROW}",
        CellIsRule(
            operator="equal",
            formula=[f'"{CHECK}"'],
            fill=fill("1E8449"),
            font=font(16, True, WHITE, TICK_FONT),
        ),
    )
    row_fill_range = f"A{FIRST_DATA_ROW}:C{LAST_DATA_ROW}"
    ws.conditional_formatting.add(
        row_fill_range,
        FormulaRule(formula=[f'$N{FIRST_DATA_ROW}="超出"'], fill=fill(RED_BG), stopIfTrue=True),
    )
    ws.conditional_formatting.add(
        row_fill_range,
        FormulaRule(formula=[f'$N{FIRST_DATA_ROW}="齊"'], fill=fill(GREEN_BG), stopIfTrue=True),
    )
    ws.conditional_formatting.add(
        row_fill_range,
        FormulaRule(formula=[f'$N{FIRST_DATA_ROW}="未齊"'], fill=fill(AMBER_BG), stopIfTrue=True),
    )

    ws.row_dimensions[TOTAL_ROW].height = 24
    total_fill = fill(NAVY)
    total_font = font(12, True, WHITE)
    merge(ws, f"A{TOTAL_ROW}:B{TOTAL_ROW}", "合計", total_font, total_fill, center, border(NAVY))
    apply(
        ws.cell(TOTAL_ROW, 3),
        f"=SUM(C{FIRST_DATA_ROW}:C{LAST_DATA_ROW})",
        total_font,
        total_fill,
        center,
        border(NAVY),
        fmt="0",
    )
    for col in range(4, 12):
        apply(ws.cell(TOTAL_ROW, col), None, total_font, total_fill, center, border(NAVY))
    apply(
        ws.cell(TOTAL_ROW, 12),
        f"=SUM(L{FIRST_DATA_ROW}:L{LAST_DATA_ROW})",
        total_font,
        total_fill,
        center,
        border(NAVY),
        fmt="0",
    )
    apply(
        ws.cell(TOTAL_ROW, 13),
        f"=SUM(M{FIRST_DATA_ROW}:M{LAST_DATA_ROW})",
        total_font,
        total_fill,
        center,
        border(NAVY),
        fmt="0",
    )
    apply(
        ws.cell(TOTAL_ROW, 14),
        f'=COUNTIF(N{FIRST_DATA_ROW}:N{LAST_DATA_ROW},"齊")&" 組齊"',
        total_font,
        total_fill,
        center,
        border(NAVY),
    )
    apply(
        ws.cell(TOTAL_ROW, 15),
        f'=IF(C{TOTAL_ROW}={EXPECTED_TOTAL},"核對正確 51 位","人數不符，請檢查")',
        total_font,
        total_fill,
        center,
        border(NAVY),
    )

    ws.print_area = f"A1:O{TOTAL_ROW}"
    ws.page_setup.horizontalCentered = True
    ws.sheet_view.selection[0].activeCell = "B3"
    ws.sheet_view.selection[0].sqref = "B3"

    ws.protection.enable()
    ws.protection.selectLockedCells = False
    ws.protection.selectUnlockedCells = False
    # Let them widen a column on the day if a note is long.
    ws.protection.formatColumns = False
    ws.protection.formatRows = False


def build_extras(ws):
    ws.sheet_properties.tabColor = GOLD_HEADER
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 130
    ws.page_setup.orientation = "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.horizontalCentered = True
    ws.print_options.horizontalCentered = True
    ws.page_margins = PageMargins(left=0.6, right=0.6, top=0.5, bottom=0.5, header=0.2, footer=0.2)
    ws.oddFooter.center.text = "船河樂2026  臨時加人（不計入 51 位）"
    ws.oddFooter.center.font = FONT_NAME
    ws.oddFooter.center.size = 9

    widths = {"A": 8, "B": 22, "C": 10, "D": 10, "E": 18, "F": 28}
    for col, width in widths.items():
        ws.column_dimensions[col].width = width

    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left = Alignment(horizontal="left", vertical="center", indent=1)

    ws.row_dimensions[1].height = 32
    ws.row_dimensions[2].height = 32
    merge(
        ws, "A1:F1",
        "臨時加人",
        font(20, True, WHITE),
        fill(GOLD_HEADER),
        center,
        border(GOLD_HEADER),
    )
    merge(
        ws, "A2:F2",
        "名單以外的人寫在這裡，另行點算。不要改「點名表」的人數，以免 51 位對不上。",
        font(11, False, "6B5A1E"),
        fill(GOLD_BG),
        center,
        border("F3E3A8"),
    )

    headers = ["序", "姓名", "人數", "已到", "跟哪一組", "備註"]
    ws.row_dimensions[3].height = 22
    for col, text in enumerate(headers, 1):
        apply(ws.cell(3, col), text, font(11, True, WHITE), fill(NAVY), center, border(NAVY))

    first, last = 4, 15
    thin = border()
    for i, row in enumerate(range(first, last + 1)):
        ws.row_dimensions[row].height = 22
        band = fill(PAPER) if i % 2 else fill(WHITE)
        apply(ws.cell(row, 1), i + 1, font(12, True, NAVY), band, center, thin)
        apply(ws.cell(row, 2), None, font(12), band, left, thin, locked=False)
        apply(ws.cell(row, 3), None, font(12), band, center, thin, locked=False, fmt="0")
        apply(ws.cell(row, 4), None, font(12, True, TEAL), band, center, thin, locked=False, fmt="0")
        apply(ws.cell(row, 5), None, font(12), band, left, thin, locked=False)
        apply(ws.cell(row, 6), None, font(12), band, left, thin, locked=False)

    number_dv = DataValidation(
        type="whole",
        operator="greaterThanOrEqual",
        formula1="0",
        allow_blank=True,
        showErrorMessage=True,
        showInputMessage=False,
        errorTitle="人數",
        error="請填 0 或以上的整數",
        errorStyle="stop",
    )
    number_dv.add(f"C{first}:D{last}")
    ws.add_data_validation(number_dv)

    total_row = last + 1
    ws.row_dimensions[total_row].height = 24
    merge(ws, f"A{total_row}:B{total_row}", "臨時合計", font(12, True, WHITE), fill(NAVY), center, border(NAVY))
    apply(ws.cell(total_row, 3), f"=SUM(C{first}:C{last})", font(12, True, WHITE), fill(NAVY), center, border(NAVY), fmt="0")
    apply(ws.cell(total_row, 4), f"=SUM(D{first}:D{last})", font(12, True, WHITE), fill(NAVY), center, border(NAVY), fmt="0")
    apply(ws.cell(total_row, 5), None, font(12, True, WHITE), fill(NAVY), center, border(NAVY))
    apply(ws.cell(total_row, 6), "不計入 51 位", font(11, True, WHITE), fill(NAVY), center, border(NAVY))

    ws.print_area = f"A1:F{total_row}"
    ws.freeze_panes = "A4"
    ws.protection.enable()
    ws.protection.selectLockedCells = False
    ws.protection.selectUnlockedCells = False
    ws.protection.formatColumns = False
    ws.protection.formatRows = False


if __name__ == "__main__":
    path = build()
    print(path)
